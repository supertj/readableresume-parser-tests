"""Negative tests: the checks must fail on inputs that look fine at a glance.

Each case takes a passing fixture, breaks it in one specific way, and asserts the
relevant check now fails. Runs before every `make check`.

SPDX-License-Identifier: AGPL-3.0-only
"""

import copy
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).parent))

from compare import check_text  # noqa: E402
from docx_text import check_docx  # noqa: E402
from pdf_text import read_pdf  # noqa: E402
from verify import check_fonts  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "fixtures" / "jakes"
EXPECTED = json.loads((FIXTURES / "expected-resume.json").read_text())
PDF = FIXTURES / "jakes-resume-letter.pdf"
DOCX = FIXTURES / "jakes-resume-letter.docx"


def case(name: str, failed: bool) -> bool:
    print(f"{'ok  ' if failed else 'MISS'}  {name}")
    return failed


def broken_docx(mutate) -> dict:
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "broken.docx"
        shutil.copy(DOCX, path)
        document = Document(path)
        mutate(document)
        document.save(path)
        return check_docx(str(path), EXPECTED)


def _pdf_with_form_xobject_text(text: str) -> bytes:
    """A one-page PDF whose only text sits in a Form XObject, in unembedded Helvetica."""
    form_stream = f"BT /F1 12 Tf 72 700 Td ({text}) Tj ET".encode()
    page_stream = b"q /X0 Do Q"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /XObject << /X0 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(page_stream) + page_stream + b"\nendstream",
        b"<< /Type /XObject /Subtype /Form /BBox [0 0 612 792] /Resources << /Font << /F1 6 0 R >> >> /Length %d >>\nstream\n"
        % len(form_stream)
        + form_stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(out)


def main() -> None:
    lines = read_pdf(str(PDF)).lines
    # The unmodified fixture must pass; every case below breaks it and must fail
    baseline = check_text(EXPECTED, lines)["ok"] and check_docx(str(DOCX), EXPECTED)["ok"]
    print(f"{'ok  ' if baseline else 'FAIL'}  unmodified fixture passes")
    results = [baseline]

    results.append(case("duplicated content is caught", not check_text(EXPECTED, lines + lines)["ok"]))
    results.append(case("an extra unexpected line is caught", not check_text(EXPECTED, lines + ["Hobbies: chess"])["ok"]))
    results.append(case("a missing bullet is caught", not check_text(EXPECTED, [l for l in lines if "Grafana" not in l])["ok"]))
    swapped = copy.copy(lines)
    i, j = next(k for k, l in enumerate(swapped) if "EXPERIENCE" in l), next(k for k, l in enumerate(swapped) if "PROJECTS" in l)
    swapped[i], swapped[j] = swapped[j], swapped[i]
    results.append(case("reordered sections are caught", not check_text(EXPECTED, swapped)["ok"]))
    results.append(case("a split word is caught", not check_text(EXPECTED, [l.replace("Projects", "Proj ects").replace("PROJECTS", "PROJ ECTS") for l in lines])["ok"]))
    results.append(case("a ligature is caught", not check_text(EXPECTED, [l.replace("fi", "ﬁ") for l in lines])["ok"]))
    results.append(case("a private-use bullet is caught", not check_text(EXPECTED, [l.replace("•", "") for l in lines])["ok"]))

    def disable_numbering(document):
        for p in document.paragraphs:
            if p._p.pPr is not None and p._p.pPr.numPr is not None:
                p._p.pPr.numPr.numId.val = 0

    results.append(case("numbering set to numId 0 is caught", not broken_docx(disable_numbering)["ok"]))

    def add_table(document):
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Hidden text"

    results.append(case("a table in the .docx is caught", not broken_docx(add_table)["ok"]))

    def add_header(document):
        document.sections[0].header.paragraphs[0].text = "jordan.ellis@email.com"

    results.append(case("text in a header is caught", not broken_docx(add_header)["ok"]))

    def symbol_bullet(document):
        numbering = document.part.numbering_part.element
        for fonts in numbering.iter(qn("w:rFonts")):
            fonts.set(qn("w:ascii"), "Symbol")

    results.append(case("a bullet on the Symbol font is caught", not broken_docx(symbol_bullet)["ok"]))

    # A location that appears once can't satisfy two expected occurrences
    once = list(lines)
    first = next(i for i, l in enumerate(once) if "Chicago, IL" in l)
    once[first] = once[first].replace("Chicago, IL", "", 1)
    once_result = check_text(EXPECTED, once)
    results.append(case("a repeated value missing one occurrence is caught", once_result["missingOccurrences"] == ["Chicago, IL"]))

    def symbol_override(document):
        numbering = document.part.numbering_part.element
        num = numbering.find(qn("w:num"))
        override = num.makeelement(qn("w:lvlOverride"), {qn("w:ilvl"): "0"})
        level = override.makeelement(qn("w:lvl"), {qn("w:ilvl"): "0"})
        for tag, value in (("w:numFmt", "bullet"), ("w:lvlText", "\uf0b7")):
            level.append(level.makeelement(qn(tag), {qn("w:val"): value}))
        override.append(level)
        num.append(override)

    results.append(case("a private-use glyph in a level override is caught", not broken_docx(symbol_override)["ok"]))

    def table_in_header(document):
        header = document.sections[0].header
        header.is_linked_to_previous = False
        header.add_table(rows=1, cols=1, width=document.sections[0].page_width).cell(0, 0).text = "Hidden"

    results.append(case("text in a header table is caught", not broken_docx(table_in_header)["ok"]))

    def move_bullet(document):
        paragraphs = [p for p in document.paragraphs if p.text.strip()]
        bullet = next(p for p in paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None)
        name = paragraphs[0]
        name._p.get_or_add_pPr().append(copy.deepcopy(bullet._p.pPr.numPr))
        bullet._p.pPr.remove(bullet._p.pPr.numPr)

    results.append(case("a bullet moved to another paragraph is caught", not broken_docx(move_bullet)["ok"]))

    with tempfile.TemporaryDirectory() as tmp:
        hidden = Path(tmp) / "xobject.pdf"
        hidden.write_bytes(_pdf_with_form_xobject_text("Hidden"))
        facts = read_pdf(str(hidden))
        results.append(case("text inside a Form XObject is extracted", any("Hidden" in l for l in facts.lines)))
        results.append(case("an unembedded font inside a Form XObject is caught", not check_fonts(facts, "Arial")["ok"]))

    typo = subprocess.run([sys.executable, str(ROOT / "checks" / "verify.py"), "no-such-template"], capture_output=True)
    results.append(case("an unknown selector exits with an error", typo.returncode != 0))

    if not all(results):
        sys.exit("selftest failed: a check accepts broken input")
    print("selftest passed")


if __name__ == "__main__":
    main()
