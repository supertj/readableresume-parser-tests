"""Tests for the rules analyze.py uses to sort failed fields and to find where text sits in a .docx, and a check
that every layout in docx/ keeps the control's text formatting. Runs before every `make layout-study`.

Usage, from the repository root:

    uv run python findings/2026-10-28-layout/test_analyze.py

SPDX-License-Identifier: AGPL-3.0-only
"""

import re
import sys
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from lxml import etree

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from analyze import ORDER, docx_summary, docx_text, lost_by_group, where  # noqa: E402

TEXT_BOX_RUN = (
    '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    ' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
    ' xmlns:v="urn:schemas-microsoft-com:vml"><mc:AlternateContent>'
    "<mc:Choice Requires=\"wps\"><w:drawing><w:txbxContent><w:p><w:r><w:t>Boxed Name</w:t></w:r></w:p></w:txbxContent></w:drawing></mc:Choice>"
    "<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>Boxed Name</w:t></w:r></w:p></w:txbxContent>"
    "</v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r>"
)


def case(name: str, passed: bool) -> bool:
    print(f"{'ok  ' if passed else 'FAIL'}  {name}")
    return passed


def sample_docx(folder: Path) -> Path:
    """Header text, a body paragraph that anchors a text box (with a fallback copy), and a table."""
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Header Name"
    anchor = document.add_paragraph("Body line")
    anchor._p.append(parse_xml(TEXT_BOX_RUN))
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Left cell"
    table.cell(0, 1).text = "Right cell"
    path = folder / "sample.docx"
    document.save(str(path))
    return path


# Parts that set formatting outside the runs: style definitions and document defaults, list (bullet) formatting,
# theme fonts. Every layout must carry the control's copy byte for byte, so inherited formatting can't change either.
SHARED_FORMATTING = re.compile(r"word/(styles|stylesWithEffects|numbering|fontTable)\.xml|word/theme/theme\d*\.xml")


def text_formats(archive: zipfile.ZipFile) -> Counter:
    """(text, paragraph style, table style, run formatting) for every run with text in the body, tables, text boxes
    and headers.

    Order doesn't matter: the layouts move text around. Formatting does, because a different size or style changes
    where lines wrap and what the parser sees, and a lost field could no longer be put down to the layout alone.
    Separators and spaces are left out: the icons layout replaces the | separators with pictures.
    """
    found = Counter()
    for name in archive.namelist():
        if not re.fullmatch(r"word/(document|header\d*)\.xml", name):
            continue
        for run in etree.fromstring(archive.read(name)).iter(qn("w:r")):
            # Only the run's own text: the run that holds a text box contains the box's runs, counted separately
            text = "".join(node.text or "" for node in run.findall(qn("w:t")))
            if text.strip(" |") == "":
                continue
            style = next(run.iterancestors(qn("w:p"))).find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
            # A table style can carry its own font and size; with none, cells use the document defaults
            table = next(run.iterancestors(qn("w:tbl")), None)
            table_style = None if table is None else table.find(f"{qn('w:tblPr')}/{qn('w:tblStyle')}")
            run_format = run.find(qn("w:rPr"))
            found[
                (
                    text,
                    None if style is None else style.get(qn("w:val")),
                    None if table_style is None else table_style.get(qn("w:val")),
                    # Elements and attributes, without namespace declarations (runs in a text box carry extra ones)
                    () if run_format is None else tuple((el.tag, tuple(sorted(el.attrib.items()))) for el in run_format.iter()),
                )
            ] += 1
    return found


def formatting_changes(control: Path, layout: Path) -> list[str]:
    """What in this layout's text formatting differs from the control: shared parts first, then runs of text."""
    with zipfile.ZipFile(control) as base, zipfile.ZipFile(layout) as other:
        parts = {name for name in base.namelist() + other.namelist() if SHARED_FORMATTING.fullmatch(name)}
        changes = [
            name
            for name in sorted(parts)
            if name not in base.namelist() or name not in other.namelist() or base.read(name) != other.read(name)
        ]
        runs, expected = text_formats(other), text_formats(base)
    changes += [repr(key[0]) for key in (runs - expected) + (expected - runs)]
    return changes


def edited_copy(source: Path, target: Path, part: str, old: str, new: str) -> Path:
    """A copy of a .docx with the first match of a string replaced in one part, for the negative checks."""
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(target, "w") as out:
        for item in archive.infolist():
            data = archive.read(item)
            if item.filename == part:
                assert old.encode() in data, f"{old} not in {part}"
                data = data.replace(old.encode(), new.encode(), 1)
            out.writestr(item, data)
    return target


def main() -> None:
    control = HERE / "docx" / "single-column.docx"
    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        text = docx_text(sample_docx(folder))
        # Negative checks: a bigger default font (inherited by every run) and one run set larger must both be caught
        bigger_default = edited_copy(control, folder / "default.docx", "word/styles.xml", '<w:sz w:val="22"/>', '<w:sz w:val="36"/>')
        bigger_run = edited_copy(control, folder / "run.docx", "word/document.xml", '<w:i/><w:sz w:val="21"/>', '<w:i/><w:sz w:val="22"/>')
        caught = [formatting_changes(control, bigger_default), formatting_changes(control, bigger_run)]
    changed = {path.stem: formatting_changes(control, path) for path in sorted((HERE / "docx").glob("*.docx"))}
    results = [
        case(
            "failed fields are counted by group, groups with nothing lost are left out",
            lost_by_group(["profile.name", "profile.email", "workExperiences[1].date", "workExperiences[1].descriptions[0]", "skills.descriptions[2]"])
            == {"contact": 2, "jobs": 1, "job bullets": 1, "skills": 1},
        ),
        case("header text is filed under header, not body", where("Header Name", text) == "header" and "header name" not in text["body"]),
        case("text box text is filed under text box, not body", where("Boxed Name", text) == "text box" and "boxed name" not in text["body"]),
        case("the fallback copy of a text box is not counted again", text["text box"].count("boxed name") == 1),
        case("table cells count as body, left cell first", text["body"].index("left cell") < text["body"].index("right cell")),
        case("a value that isn't anywhere is missing", where("Not in the file", text) == "missing"),
        case(
            "contact details split between places are reported as both",
            docx_summary({"contactIn": {"name": "header", "phone": "body", "email": "body", "url": "body"}, "valuesInBody": 1,
                          "valuesAsText": 1, "valuesTotal": 1, "tables": 0, "textBoxes": 0, "images": 0})["contactIn"]
            == "body and header",
        ),
        # The committed Google Docs export of the text box layout: the text is still text in the .docx
        case(
            "Google Docs' .docx keeps the text box contents as text",
            where("jordan.ellis@example.com", docx_text(HERE / "gdocs-docx" / "text-box-contact.docx")) == "text box",
        ),
        # Each layout changes one thing: the text keeps the control's fonts, sizes, bold and italics
        case("a changed default font size is caught", caught[0] == ["word/styles.xml"]),
        case("a changed run size is caught", len(caught[1]) == 2),
        case(
            "every layout in docx/ keeps the control's text formatting",
            sorted(changed) == sorted(ORDER) and not any(changed.values()),
        ),
    ]
    if not all(results):
        sys.exit(1)


if __name__ == "__main__":
    main()
